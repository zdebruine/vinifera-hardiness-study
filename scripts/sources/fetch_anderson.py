#!/usr/bin/env python3
"""Fill variety_region.bearing_ha from Anderson & Nelgen regional bearing-area data.

Anderson, K. & Nelgen, S. (2020), *Which Winegrape Varieties are Grown Where?* (Univ. of
Adelaide). The dataset ships as Excel workbooks (region x variety bearing hectares, with a
synonym crosswalk). Free to use -- cite, do not redistribute the raw file.

Split by testability:
  * ``download`` -- NETWORK; fetches a workbook to data/raw/ (URL/filename to confirm on
    first run; not redistributed). Runs on the open-network side.
  * ``extract_hectares`` -- reads a local .xlsx (openpyxl) into (region, variety, ha) tuples.
  * ``match_to_db`` -- PURE; maps Anderson region/variety names to our ids via synonyms and
    normalized names, returning bearing_ha to write back. Unit-tested offline.

Because Anderson's region granularity differs from ours, matching is name-based and
conservative: only confident matches are written; the rest are reported for manual review.
"""
from __future__ import annotations

import csv
import sys
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts" / "sources"))
from fetch_fps_catalog import normalize  # reuse the normalizer

VARIETY_REGION = ROOT / "data" / "db" / "variety_region.csv"
VARIETIES = ROOT / "data" / "db" / "varieties.csv"
SYNONYMS = ROOT / "data" / "db" / "synonyms.csv"
RAW_DIR = ROOT / "data" / "raw"

# Confirm exact workbook URL/filename against the Adelaide publication on first run.
ANDERSON_HINT = "https://www.adelaide.edu.au/press/titles/winegrapes (database workbooks)"


def download(url: str, dest_name: str) -> Path:
    """NETWORK: fetch an Anderson workbook into data/raw/ (kept local, not committed)."""
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    dest = RAW_DIR / dest_name
    with urllib.request.urlopen(url, timeout=120) as resp:
        dest.write_bytes(resp.read())
    return dest


def extract_hectares(xlsx_path: Path, sheet: str | None = None) -> list[tuple[str, str, float]]:
    """Read a region x variety bearing-hectares sheet into (region, variety, ha) rows.

    Assumes the first column is the variety and remaining columns are regions (the common
    Anderson layout). Confirm/parametrize the sheet name and orientation on first run.
    """
    from openpyxl import load_workbook  # network-free, but optional dep
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = rows[0]
    regions = [str(h) for h in header[1:]]
    out: list[tuple[str, str, float]] = []
    for r in rows[1:]:
        variety = r[0]
        if variety is None:
            continue
        for region, val in zip(regions, r[1:]):
            if isinstance(val, (int, float)) and val > 0:
                out.append((str(region), str(variety), float(val)))
    return out


def _build_index() -> tuple[dict[str, str], dict[str, str]]:
    """Return (normalized variety name/synonym -> variety_id), (normalized region name -> region_id)."""
    var_index: dict[str, str] = {}
    with VARIETIES.open(newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            var_index[normalize(r["prime_name"])] = r["variety_id"]
    with SYNONYMS.open(newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            var_index.setdefault(normalize(r["synonym"]), r["variety_id"])
    # region matching is left to the caller's own region name map; ours is name-based.
    return var_index, {}


def match_to_db(anderson_rows: list[tuple[str, str, float]],
                var_index: dict[str, str],
                region_name_to_id: dict[str, str]) -> dict[tuple[str, str], float]:
    """PURE: ((region_id, variety_id) -> bearing_ha) for confident name matches only."""
    out: dict[tuple[str, str], float] = {}
    for region, variety, ha in anderson_rows:
        vid = var_index.get(normalize(variety))
        rid = region_name_to_id.get(normalize(region))
        if vid and rid:
            out[(rid, vid)] = out.get((rid, vid), 0.0) + ha
    return out


def apply_bearing_ha(matches: dict[tuple[str, str], float]) -> int:
    """Write matched hectares into variety_region.csv (only where the pair already exists)."""
    with VARIETY_REGION.open(newline="", encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        cols = reader.fieldnames
        rows = list(reader)
    n = 0
    for row in rows:
        key = (row["region_id"], row["variety_id"])
        if key in matches:
            row["bearing_ha"] = f"{matches[key]:.0f}"
            n += 1
    with VARIETY_REGION.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=cols)
        writer.writeheader()
        writer.writerows(rows)
    return n


if __name__ == "__main__":
    print("Anderson loader. Provide a downloaded workbook path; region-name mapping to our "
          "region_ids must be supplied (granularity differs).")
    print(f"Source: {ANDERSON_HINT}")
